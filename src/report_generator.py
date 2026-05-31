from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, PatternFill


HEADER_FILL = "D9EAD3"
WARNING_FILL = "FFF2CC"


def prepare_detail_dataframe(consolidado):
    columns = [
        "id_cobranca",
        "paciente",
        "cpf_beneficiario",
        "registro_ans",
        "ans",
        "nome_operadora",
        "procedimento",
        "descricao_servico",
        "cod_tuss",
        "data_atendimento",
        "dt_realizacao",
        "dt_lancamento",
        "valor",
        "vl_servico",
        "vl_glosa",
        "vl_liquido",
        "divergencias",
    ]

    selected_columns = [
        col for col in columns
        if col in consolidado.columns
    ]

    detail = consolidado[selected_columns].copy()

    detail = detail.rename(
        columns={
            "id_cobranca": "ID Cobrança",
            "paciente": "Paciente Excel",
            "cpf_beneficiario": "CPF",
            "registro_ans": "ANS Excel",
            "ans": "ANS CSV",
            "nome_operadora": "Convênio",
            "procedimento": "Procedimento Excel",
            "descricao_servico": "Procedimento CSV",
            "cod_tuss": "Código TUSS",
            "data_atendimento": "Data Atendimento Excel",
            "dt_realizacao": "Data Realização CSV",
            "dt_lancamento": "Data Lançamento",
            "valor": "Valor Excel",
            "vl_servico": "Valor Bruto CSV",
            "vl_glosa": "Valor Glosa",
            "vl_liquido": "Valor Líquido CSV",
            "divergencias": "Divergências",
        }
    )

    return detail


def prepare_alerts_dataframe(detail):
    alert_columns = [
        "ID Cobrança",
        "Paciente Excel",
        "CPF",
        "Convênio",
        "Procedimento Excel",
        "Valor Excel",
        "Valor Líquido CSV",
        "Valor Glosa",
        "Divergências",
    ]

    alertas = detail[
        detail["Divergências"] != "SEM_DIVERGENCIAS"
    ].copy()

    selected_columns = [
        col for col in alert_columns
        if col in alertas.columns
    ]

    return alertas[selected_columns]


def highlight_divergence_column(sheet):
    divergence_column = None

    for cell in sheet[1]:
        if cell.value == "Divergências":
            divergence_column = cell.column
            break

    if not divergence_column:
        return

    for row in sheet.iter_rows(min_row=2):
        divergence_cell = row[divergence_column - 1]

        if (
            divergence_cell.value
            and divergence_cell.value != "SEM_DIVERGENCIAS"
        ):
            divergence_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=WARNING_FILL
            )
            divergence_cell.font = Font(bold=True)


def apply_formatting(workbook):
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=HEADER_FILL
            )

        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            sheet.column_dimensions[
                column_letter
            ].width = max_length + 2

    if "Detalhamento" in workbook.sheetnames:
        highlight_divergence_column(
            workbook["Detalhamento"]
        )

    if "Alertas" in workbook.sheetnames:
        highlight_divergence_column(
            workbook["Alertas"]
        )


def generate_report(consolidado, pdf_summary=None):
    output_dir = Path("output/relatorios")
    output_dir.mkdir(parents=True, exist_ok=True)

    period = datetime.now().strftime("%Y%m")
    report_path = output_dir / f"relatorio_faturamento_{period}.xlsx"

    detail = prepare_detail_dataframe(consolidado)
    alertas = prepare_alerts_dataframe(detail)

    resumo_data = {
        "Indicador": [
            "Total de cobranças consolidadas",
            "Presentes nas duas fontes",
            "Apenas no CSV",
            "Apenas no Excel",
            "Total de alertas",
            "Valor líquido total",
            "Total de glosas",
        ],
        "Valor": [
            len(consolidado),
            int((consolidado["_merge"] == "both").sum()),
            int((consolidado["_merge"] == "right_only").sum()),
            int((consolidado["_merge"] == "left_only").sum()),
            len(alertas),
            consolidado["vl_liquido_normalizado"].sum(),
            consolidado["vl_glosa"]
            .astype(str)
            .str.replace(",", ".")
            .astype(float)
            .sum(),
        ],
    }

    resumo = pd.DataFrame(resumo_data)

    if pdf_summary:
        pdf_resumo = pd.DataFrame(
            {
                "Indicador": [
                    "PDFs processados",
                    "PDFs renomeados",
                    "PDFs para revisão manual",
                ],
                "Valor": [
                    pdf_summary.get("total", 0),
                    pdf_summary.get("renomeados", 0),
                    pdf_summary.get("revisao_manual", 0),
                ],
            }
        )

        resumo = pd.concat(
            [resumo, pdf_resumo],
            ignore_index=True
        )

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        resumo.to_excel(
            writer,
            sheet_name="Resumo",
            index=False
        )

        detail.to_excel(
            writer,
            sheet_name="Detalhamento",
            index=False
        )

        alertas.to_excel(
            writer,
            sheet_name="Alertas",
            index=False
        )

        workbook = writer.book
        apply_formatting(workbook)

    return report_path